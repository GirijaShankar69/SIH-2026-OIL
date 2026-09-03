"""
Unit tests for the SIH 2026 Presentation Generator (Baghewala Field Digital Twin).
"""

import os
import unittest

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


class TestPresentationGeneration(unittest.TestCase):
    """Test suite for SIH2026-Baghewala-Digital-Twin.pptx structure and content."""

    @classmethod
    def setUpClass(cls):
        if not HAS_PPTX:
            raise unittest.SkipTest("python-pptx is not installed")
        cls.project_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        cls.pptx_path = os.path.join(cls.project_dir, "SIH2026-Baghewala-Digital-Twin.pptx")
        
        # Ensure presentation exists, if not generate it
        if not os.path.exists(cls.pptx_path):
            import generate_presentation
            generate_presentation.main()
            
        cls.prs = Presentation(cls.pptx_path)

    def test_file_exists_and_non_empty(self):
        """Verify PPTX file exists and has substantial content."""
        self.assertTrue(os.path.exists(self.pptx_path))
        file_size = os.path.getsize(self.pptx_path)
        self.assertGreater(file_size, 20_000, "Presentation file is unexpectedly small")

    def test_slide_count_exactly_six(self):
        """Verify the presentation has exactly 6 slides matching SIH rules."""
        self.assertEqual(len(self.prs.slides), 6, "SIH Hackathon mandates exactly 6 slides")

    def test_widescreen_dimensions(self):
        """Verify 16:9 widescreen layout (13.333 x 7.5 inches)."""
        width = self.prs.slide_width.inches
        height = self.prs.slide_height.inches
        self.assertAlmostEqual(width, 13.333, places=2)
        self.assertAlmostEqual(height, 7.500, places=2)

    def test_slide_1_title_and_metadata(self):
        """Verify Slide 1 contains required problem statement metadata."""
        slide1 = self.prs.slides[0]
        text = " ".join([p.text for s in slide1.shapes if s.has_text_frame for p in s.text_frame.paragraphs])
        self.assertIn("SIH-2026-OIL-01", text)
        self.assertIn("SIH-TEAM-BGW", text)
        self.assertIn("Oil India Limited", text)
        self.assertIn("Software", text)
        self.assertIn("Baghewala", text)

    def test_slide_2_problem_and_solution_cards(self):
        """Verify Slide 2 contains the 3-column problem/solution/moat breakdown."""
        slide2 = self.prs.slides[1]
        text = " ".join([p.text for s in slide2.shapes if s.has_text_frame for p in s.text_frame.paragraphs])
        self.assertIn("Viscosity", text)
        self.assertIn("Rod Floating", text)
        self.assertIn("Boberg-Lantz", text)
        self.assertIn("VFD", text)
        self.assertIn("Classifier", text)

    def test_slide_3_five_layer_architecture_and_shapes(self):
        """Verify Slide 3 contains the 5 layers and flow elements."""
        slide3 = self.prs.slides[2]
        text = " ".join([p.text for s in slide3.shapes if s.has_text_frame for p in s.text_frame.paragraphs])
        self.assertIn("THERMAL RESERVOIR", text)
        self.assertIn("WELLBORE HYDRAULICS", text)
        self.assertIn("SRP WAVE MECHANICS", text)
        self.assertIn("AI & OPTIMIZER", text)
        self.assertIn("TWIN & VFD CONTROL", text)
        # Check shape count includes cards + arrows
        self.assertGreaterEqual(len(slide3.shapes), 20, "Slide 3 must contain layer boxes, badges and connectors")

    def test_slide_4_feasibility_and_roadmap(self):
        """Verify Slide 4 contains 4 quadrants covering SCADA, viability, roadmap, and risk."""
        slide4 = self.prs.slides[3]
        text = " ".join([p.text for s in slide4.shapes if s.has_text_frame for p in s.text_frame.paragraphs]).lower()
        self.assertIn("scada", text)
        self.assertIn("opc-ua", text)
        self.assertIn("roadmap", text)
        self.assertIn("risk", text)
        self.assertIn("phase 1", text)


    def test_slide_5_hero_metrics_and_roi(self):
        """Verify Slide 5 contains all verified hero metrics and ESG benefits."""
        slide5 = self.prs.slides[4]
        text = " ".join([p.text for s in slide5.shapes if s.has_text_frame for p in s.text_frame.paragraphs])
        self.assertIn("+16.8%", text)
        self.assertIn("-24.5%", text)
        self.assertIn("-22.3%", text)
        self.assertIn("38.2", text)
        self.assertIn("100%", text)
        self.assertIn("5.73", text)
        self.assertIn("CO₂", text)

    def test_slide_6_all_sixteen_references(self):
        """Verify Slide 6 contains all 16 academic references."""
        slide6 = self.prs.slides[5]
        text = " ".join([p.text for s in slide6.shapes if s.has_text_frame for p in s.text_frame.paragraphs])
        for i in range(1, 17):
            self.assertIn(f"{i}.", text, f"Reference {i} missing from Slide 6")
        self.assertIn("Marx", text)
        self.assertIn("Boberg", text)
        self.assertIn("Gibbs", text)
        self.assertIn("Ramey", text)
        self.assertIn("Miner", text)
        self.assertIn("Goodman", text)
        self.assertIn("Breiman", text)


if __name__ == "__main__":
    unittest.main()
